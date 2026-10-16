"""Moteur repris de CV_generator/app.py, sans son interface historique."""

import os

import json

import re

import difflib

from io import BytesIO

from datetime import datetime

import streamlit as st

from docx import Document

from docx.shared import Pt, Cm, RGBColor

from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx.oxml.ns import qn

from google import genai

from google.genai import types

from google.genai.errors import APIError

try:
    from openai import OpenAI
except ImportError:
    OpenAI = None

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.enums import TA_RIGHT, TA_JUSTIFY
    from reportlab.lib.colors import HexColor
    REPORTLAB_OK = True
except ImportError:
    REPORTLAB_OK = False

COLOR_PALETTE = {
    "Bleu Roi": {"hex": "#4169E1", "hover": "#3151B5", "rgb": RGBColor(65, 105, 225)},
    "Rose":     {"hex": "#E91E63", "hover": "#C2185B", "rgb": RGBColor(233, 30, 99)},
    "Vert":     {"hex": "#10B981", "hover": "#059669", "rgb": RGBColor(16, 185, 129)},
    "Rouge":    {"hex": "#EF4444", "hover": "#DC2626", "rgb": RGBColor(239, 68, 68)}
}

FONT_FAMILY = "Montserrat"

MOIS_FR = [
    "janvier", "février", "mars", "avril", "mai", "juin", 
    "juillet", "août", "septembre", "octobre", "novembre", "décembre"
]

from pathlib import Path
_profile_path = Path(os.getenv("JOBFLOW_PROFILE_PATH", str(Path(__file__).parent / "data" / "profile.json")))
if not _profile_path.is_file():
    _profile_path = Path(__file__).parent / "profile.example.json"
MASTER_CV = json.loads(_profile_path.read_text())

CV_STYLES = [
    "Chronologique (classique)",
    "Thématique (compétences)",
    "Mixte (parcours + compétences)"
]

SYSTEM_PROMPT = f"""
Tu es un expert en recrutement et en optimisation de CV pour passer les filtres ATS (Applicant Tracking Systems).
RÈGLE D'OR ABSOLUE : LA VÉRITÉ STRICTE.
Tu ne dois JAMAIS inventer d'expérience, d'employeur, de rôle, de date, de diplôme, de certification, d'outil, de langue ou de chiffre.
Le parcours d'Candidate est figé dans le CV maître suivant :
{json.dumps(MASTER_CV, ensure_ascii=False, indent=2)}

RÈGLES DE COMPATIBILITÉ ATS & STYLE (CV) :
1. Intègre mot pour mot les mots-clés de l'offre.
2. RÈGLE DES VERBES : Utilise TOUJOURS des verbes à L'INFINITIF en début de chaque puce.
3. Ne modifie pas la structure de base des faits du CV.
4. Inclus impérativement les sections `languages` et `volunteering`.
5. Rédige l'ensemble au FÉMININ pour Candidate.

RÈGLES POUR LE CV (inspirées du guide Pôle Emploi) :
- Le CV s'ouvre sur une ACCROCHE professionnelle combinant « Objectif » + « Compétences clés ».
- Le titre du CV doit être un TITRE SEO optimisé : précis, mots-clés métier, 3 à 6 mots maximum.
  Exemple : « Ingénieure PMO & Cheffe de Projet IT » plutôt que « Ingénieure ».
- La rubrique des centres d'intérêt doit être nommée « MAIS AUSSI… » (plus convivial).
- Selon le modèle demandé (chronologique / thématique / mixte), adapter la structure.

MODÈLE CHRONOLOGIQUE :
- Expériences de la plus récente à la plus ancienne.
- Ordre : Accroche → Compétences → Expériences → Formations → Langues → Mais aussi…

MODÈLE THÉMATIQUE :
- Regrouper les compétences en thèmes (ex. « Gestion de projet », « Data & BI », « Relation client »).
- Pour chaque thème, apporter 1 à 2 preuves issues du CV maître.
- Placer « Compétences & Réalisations » AVANT les expériences.

MODÈLE MIXTE :
- Sous chaque expérience, ajouter des sous-titres de compétences (ex. « Gestion de projet : », « Data & BI : »).
- Mettre en avant la polyvalence et les compétences transférables.

RÈGLES DE RÉDACTION DE LA LETTRE DE MOTIVATION (structure classique française) :
La lettre doit suivre EXACTEMENT la logique argumentaire « VOUS – MOI – NOUS » :
- `recipient` : bloc destinataire (formule d'appel + éventuellement nom, rôle, entreprise, adresse).
- `object` : objet, sans écrire « Lettre de motivation ». Format : « Candidature au poste de X ».
- `accroche` : 2 à 3 lignes percutantes, chiffrées si possible. Ne PAS répéter l'objet.
- `vous` : paragraphe sur l'ENTREPRISE (réputation, culture, actualité, enjeux).
- `moi` : paragraphe sur le PROFIL + 1 à 2 réalisations chiffrées. Utiliser « j'ai réalisé », « j'ai piloté ». Éviter « je pense », « je crois ».
- `nous` : paragraphe de PROMESSE de collaboration.
- `action` : paragraphe de disponibilité / demande d'entretien.
- `politesse` : formule de politesse classique et respectueuse.
- `signature` : « Candidate ».

INTERDITS dans la lettre :
- « Bonjour », « Cher Monsieur », « Chère Madame ».
- Chômage prolongé, prétentions salariales, critiques d'anciens employeurs.
- « émerveillée », « impatientée », « extrêmement chanceuse ».
- Répéter l'objet dans l'accroche.
- « Cordialement », « Bien à vous ».

Rédige au FÉMININ (Candidate).
"""

def parse_ai_json(raw_text: str) -> dict:
    cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw_text.strip(), flags=re.I | re.S)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", cleaned, re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise ValueError("Impossible d'extraire un JSON valide de la réponse de l'IA.")

def cover_letter_to_paragraphs(cl: dict) -> list:
    """Reconstruit la liste de paragraphes à partir des champs structurés."""
    if not isinstance(cl, dict):
        return []
    paragraphs = []
    greeting = (cl.get("recipient", {}) or {}).get("title") or "Madame, Monsieur,"
    paragraphs.append(greeting)
    for key in ("accroche", "vous", "moi", "nous", "action", "politesse"):
        val = cl.get(key, "")
        if val:
            paragraphs.append(val)
    return paragraphs

def call_gemini(api_key: str, prompt: str) -> str:
    if not api_key or not api_key.strip():
        raise ValueError("Clé API Gemini manquante.")
    client = genai.Client(api_key=api_key.strip())
    model_name = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")
    config = types.GenerateContentConfig(
        temperature=0.1,
        response_mime_type="application/json"
    )
    try:
        response = client.models.generate_content(
            model=model_name,
            contents=prompt,
            config=config
        )
        if not response.text:
            raise ValueError("L'API Gemini a retourné une réponse vide.")
        return response.text
    except APIError as e:
        raise RuntimeError(f"Erreur API Gemini ({e.code}): {e.message}")
    except Exception as e:
        raise RuntimeError(f"Erreur lors de la communication avec l'IA: {str(e)}")

def call_deepseek(api_key: str, prompt: str) -> str:
    if not api_key or not api_key.strip():
        raise ValueError("Clé API DeepSeek manquante.")
    if OpenAI is None:
        raise ImportError("Le package 'openai' doit être installé (`pip install openai`).")
    try:
        client = OpenAI(api_key=api_key.strip(), base_url="https://api.deepseek.com")
        response = client.chat.completions.create(
            model="deepseek-chat",
            messages=[
                {"role": "system", "content": "Tu es un assistant expert rédigeant uniquement au format JSON valide."},
                {"role": "user", "content": prompt}
            ],
            response_format={"type": "json_object"},
            temperature=0.1
        )
        return response.choices[0].message.content
    except Exception as e:
        raise RuntimeError(f"Erreur API DeepSeek : {str(e)}")

def generate_complete_application(
    job_offer_text: str,
    provider: str,
    api_key: str,
    company_context: str = "",
    recent_news: str = "",
    cv_style: str = "Chronologique (classique)",
) -> dict:
    """Génère toute la candidature avec un seul appel à l'API choisie."""
    company_context = (company_context or "").strip()
    recent_news = (recent_news or "").strip()

    prompt = f"""
{SYSTEM_PROMPT}

MODÈLE DE CV DEMANDÉ : {cv_style}

MISSION :
À partir de l'offre ci-dessous, réalise en une seule fois :
1. l'analyse structurée de l'offre ;
2. le matching factuel avec le CV maître ;
3. la stratégie ATS ;
4. le CV personnalisé (modèle : {cv_style}) ;
5. la lettre de motivation (structure « Vous – Moi – Nous »).

CONTEXTE ENTREPRISE :
{company_context if company_context else "Aucun contexte spécifique fourni."}

ACTUALITÉ RÉCENTE DE L'ENTREPRISE À INTÉGRER :
{recent_news if recent_news else "Aucune actualité récente fournie. Rester générique sur les enjeux actuels sans inventer de faits."}

OFFRE D'EMPLOI :
\"\"\"
{job_offer_text}
\"\"\"

Consignes prioritaires pour la lettre de motivation :
- Respecter strictement la structure « Vous – Moi – Nous ».
- Intégrer le contexte et l'actualité de l'entreprise sans inventer de faits.
- Ton professionnel, crédible, adapté au poste.

Retourne UNIQUEMENT un objet JSON valide respectant exactement cette structure :
{{
  "analysis": {{
    "job_title": "Titre du poste au féminin",
    "company": "Nom de l'entreprise",
    "sector": "Secteur d'activité",
    "location": "Ville ou région",
    "main_responsibilities": ["mission 1"],
    "required_skills": [{{"skill": "nom", "priority": 5}}],
    "tools": ["outil 1"],
    "ats_keywords": ["mot-clé 1"]
  }},
  "matching": {{
    "matches": [{{
      "requirement": "Exigence",
      "match": "MATCH_FORT|MATCH_PARTIEL|ABSENT",
      "evidence": "Preuve exacte issue du CV maître, ou explication si absente"
    }}]
  }},
  "strategy": {{
    "cv_title": "Titre ciblé au féminin",
    "tailored_summary": "Accroche au féminin",
    "priority_skills": ["compétence 1"],
    "highlighted_tools": ["outil 1"],
    "rationale": ["explication 1"]
  }},
  "cv": {{
    "cv_style": "chronologique|thematique|mixte",
    "seo_title": "Titre optimisé SEO (3-6 mots)",
    "accroche_objective": "Objectif professionnel + compétences clés (2-3 lignes)",
    "tailored_summary": "Accroche courte au féminin",
    "skills": ["compétence 1"],
    "skills_grouped": [
      {{ "theme": "Thème 1", "evidence": ["Preuve 1", "Preuve 2"] }}
    ],
    "tools": ["outil 1"],
    "languages": ["Français (Langue maternelle)", "Anglais professionnel (niveau B2 / Cambridge B2 First)"],
    "experiences": [{{
      "role": "Rôle au féminin",
      "company": "Entreprise",
      "period": "Dates inchangées",
      "bullets": ["Puce commençant par un verbe à l'infinitif"],
      "skill_subheadings": [
        {{ "label": "Gestion de projet :", "text": "Description courte" }}
      ]
    }}],
    "volunteering": [{{
      "role": "Rôle au féminin",
      "organization": "Association",
      "period": "Dates inchangées",
      "bullets": ["Puce commençant par un verbe à l'infinitif"]
    }}],
    "education": [{{"title": "Diplôme", "institution": "École", "year": "Année"}}],
    "interests": ["Engagement associatif", "Sport collectif"]
  }},
  "cover_letter": {{
    "recipient": {{
      "title": "Madame, Monsieur,",
      "name": "Nom du recruteur si connu, sinon chaîne vide",
      "role": "Fonction du recruteur si connue, sinon chaîne vide",
      "company": "Nom de l'entreprise",
      "address": "Adresse si connue, sinon chaîne vide"
    }},
    "object": "Candidature au poste de ...",
    "accroche": "2 à 3 lignes percutantes avec si possible un chiffre",
    "vous": "Paragraphe sur l'entreprise",
    "moi": "Paragraphe sur le profil + réalisations chiffrées",
    "nous": "Paragraphe de promesse de collaboration",
    "action": "Paragraphe de disponibilité et demande d'entretien",
    "politesse": "Je vous prie d'agréer, Madame, Monsieur, l'expression de mes salutations distinguées.",
    "signature": "{MASTER_CV['candidate_name']}"
  }}
}}

CONTRÔLE FINAL AVANT DE RÉPONDRE :
- Ne rien inventer, conserver toutes les entreprises, dates et formations du CV maître.
- Utiliser uniquement des compétences justifiables par le CV maître.
- Commencer chaque puce par un verbe à l'infinitif.
- Inclure toutes les expériences, les deux engagements bénévoles, les formations et les langues.
- Respecter la structure de la lettre (recipient, object, accroche, vous, moi, nous, action, politesse, signature).
- Répondre uniquement avec le JSON, sans commentaire ni bloc Markdown.
"""
    if provider == "Gemini":
        raw_text = call_gemini(api_key, prompt)
    else:
        raw_text = call_deepseek(api_key, prompt)

    result = parse_ai_json(raw_text)

    required_sections = {"analysis", "matching", "strategy", "cv", "cover_letter"}
    missing = required_sections.difference(result)
    if missing:
        raise ValueError("Réponse IA incomplète. Sections manquantes : " + ", ".join(sorted(missing)))

    cl = result["cover_letter"]
    required_cl_keys = {"recipient", "object", "accroche", "vous", "moi", "nous", "action", "politesse"}
    missing_cl = required_cl_keys.difference(cl)
    if missing_cl:
        raise ValueError("Lettre de motivation incomplète. Champs manquants : " + ", ".join(sorted(missing_cl)))

    cl["paragraphs"] = cover_letter_to_paragraphs(cl)
    return result

def cv_to_text(cv_dict: dict) -> str:
    text = f"{cv_dict.get('cv_title', '')}\n{cv_dict.get('tailored_summary', '')}\n"
    text += " ".join(cv_dict.get('skills', [])) + "\n"
    text += " ".join(cv_dict.get('tools', [])) + "\n"
    for exp in cv_dict.get('experiences', []):
        bullets = exp.get('bullets', [])
        b_str = " ".join(bullets) if isinstance(bullets, list) else str(bullets)
        text += f"{exp.get('role', '')} {exp.get('company', '')} {b_str}\n"
    return text

def calculate_similarity(text1: str, text2: str) -> float:
    if not text1 or not text2:
        return 0.0
    return round(difflib.SequenceMatcher(None, text1, text2).ratio() * 100, 1)

def calculate_scores(analysis: dict, matching: dict, current_cv: dict) -> dict:
    reqs = matching.get("matches", [])
    total_reqs = len(reqs)
    strong = sum(1 for m in reqs if m.get("match") == "MATCH_FORT")
    partial = sum(1 for m in reqs if m.get("match") == "MATCH_PARTIEL")

    match_score = round(((strong * 1.0 + partial * 0.5) / max(total_reqs, 1)) * 100, 1)

    ats_kws = analysis.get("ats_keywords", [])
    cv_full_text = cv_to_text(current_cv).lower()
    covered_kws = [kw for kw in ats_kws if kw.lower() in cv_full_text]
    ats_coverage = round((len(covered_kws) / max(len(ats_kws), 1)) * 100, 1)

    master_text = cv_to_text(MASTER_CV)
    curr_text = cv_to_text(current_cv)
    sim_master = calculate_similarity(curr_text, master_text)

    return {
        "match_score": match_score,
        "ats_coverage": ats_coverage,
        "sim_master": sim_master,
        "covered_kws": covered_kws,
        "uncovered_kws": [kw for kw in ats_kws if kw not in covered_kws]
    }

def normalize_fact_text(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()

def collect_master_fact_set():
    facts = set()
    facts.update(MASTER_CV.get("candidate_name", "").lower().split())
    for exp in MASTER_CV.get("experiences", []):
        facts.add(normalize_fact_text(exp.get("role", "")))
        facts.add(normalize_fact_text(exp.get("company", "")))
    for vol in MASTER_CV.get("volunteering", []):
        facts.add(normalize_fact_text(vol.get("role", "")))
        facts.add(normalize_fact_text(vol.get("organization", "")))
    for edu in MASTER_CV.get("education", []):
        facts.add(normalize_fact_text(edu.get("title", "")))
        facts.add(normalize_fact_text(edu.get("institution", "")))
    for skill in MASTER_CV.get("skills", []):
        facts.add(normalize_fact_text(skill))
    for tool in MASTER_CV.get("tools", []):
        facts.add(normalize_fact_text(tool))
    return {p for p in facts if p}

def find_numeric_anomalies(text):
    numbers = re.findall(r"\b\d+(?:[.,]\d+)?%?\b|\b\d+h\d{2}\b|\b\d+\s*h\s*\d{2}\b", text, flags=re.I)
    return [n.strip() for n in numbers]

def audit_candidature(job_offer_text: str, cv_data: dict, cover_letter_paragraphs: list, company_name: str = "", location: str = "") -> dict:
    issues = []
    critical = []
    warnings = []
    score = 100

    job_text = (job_offer_text or "").lower()
    letter_text = " ".join(cover_letter_paragraphs or []).lower()
    cv_text = cv_to_text(cv_data).lower()
    company_name_norm = normalize_fact_text(company_name)
    location_norm = normalize_fact_text(location)

    master_companies = {normalize_fact_text(exp.get("company", "")) for exp in MASTER_CV.get("experiences", [])}
    master_companies.update({normalize_fact_text(vol.get("organization", "")) for vol in MASTER_CV.get("volunteering", [])})
    master_companies.update({normalize_fact_text(edu.get("institution", "")) for edu in MASTER_CV.get("education", [])})

    master_roles = {normalize_fact_text(e.get("role", "")) for e in MASTER_CV.get("experiences", [])}

    for exp in cv_data.get("experiences", []):
        company = normalize_fact_text(exp.get("company", ""))
        role = normalize_fact_text(exp.get("role", ""))
        if company and company not in master_companies and company_name_norm and company not in company_name_norm:
            issues.append(f"Entreprise non conforme au CV maître : {exp.get('company')}")
            critical.append("Entreprise non traceable dans le CV source")
            score -= 17
        if role and role not in master_roles:
            issues.append(f"Rôle potentiellement inventé : {exp.get('role')}")
            critical.append("Rôle potentiellement inventé")
            score -= 18
        if exp.get("period"):
            dates = re.findall(r"\b(19|20)\d{2}\b", str(exp.get("period", "")))
            if dates and not any(str(y) in str(MASTER_CV).lower() for y in dates):
                warnings.append(f"Vérifier la période de {exp.get('role')} : {exp.get('period')}")
                score -= 6

    master_edu_titles = {normalize_fact_text(item.get("title", "")) for item in MASTER_CV.get("education", [])}
    for edu in cv_data.get("education", []):
        if edu.get("title") and normalize_fact_text(edu.get("title")) not in master_edu_titles:
            issues.append(f"Formation potentiellement inventée : {edu.get('title')}")
            critical.append("Formation potentiellement inventée")
            score -= 12

    allowed_skills = {normalize_fact_text(s) for s in MASTER_CV.get("skills", [])}
    allowed_tools = {normalize_fact_text(t) for t in MASTER_CV.get("tools", [])}
    for skill in cv_data.get("skills", []):
        ns = normalize_fact_text(skill)
        if ns and ns not in allowed_skills and ns not in {"pmo", "gestion de projet", "gestion de projet it"}:
            issues.append(f"Compétence suspecte : {skill}")
            warnings.append(f"Vérifier la compétence {skill}")
            score -= 8
    for tool in cv_data.get("tools", []):
        nt = normalize_fact_text(tool)
        if nt and nt not in allowed_tools:
            issues.append(f"Outil potentiellement inventé : {tool}")
            warnings.append(f"Contrôler l’outil {tool}")
            score -= 8

    numeric_issues = find_numeric_anomalies(letter_text + " " + cv_text)
    allowed_numeric = set(find_numeric_anomalies(str(MASTER_CV)))
    for number in numeric_issues:
        if number and number not in allowed_numeric and not any(number in part for part in ["2025", "2024", "2023", "2026", "2022", "2020", "2019", "2014"]) and not number.endswith("%"):
            warnings.append(f"Chiffre ou métrique à vérifier : {number}")
            score -= 5

    missing_keywords = []
    for key in re.findall(r"[A-Za-zÀ-ÖØ-öø-ÿ]+(?:\s+[A-Za-zÀ-ÖØ-öø-ÿ]+){0,2}", job_text):
        clean_key = normalize_fact_text(key)
        if len(clean_key) > 4 and clean_key not in {"poste", "projet", "travail", "client", "entreprise", "mission", "rôle"}:
            if clean_key not in cv_text and clean_key not in letter_text:
                missing_keywords.append(key)
    if missing_keywords:
        warnings.extend([f"Mot-clé important absent : {m}" for m in missing_keywords[:3]])
        score -= min(10, 3 * len(missing_keywords[:3]))

    if company_name_norm and company_name_norm not in letter_text and company_name_norm not in cv_text:
        warnings.append("La lettre ne mentionne pas explicitement l’entreprise ciblée")
        score -= 8
    if location_norm and location_norm not in letter_text and location_norm not in cv_text:
        warnings.append("Le lieu n’apparaît pas dans le message de candidature")
        score -= 4
    if "bonjour" in letter_text or "cher monsieur" in letter_text or "chère madame" in letter_text:
        warnings.append("Formule d'ouverture familière détectée (à éviter en France)")
        score -= 6
    if "cordialement" in letter_text or "bien à vous" in letter_text:
        warnings.append("Formule de politesse trop familière (à éviter)")
        score -= 5

    if re.search(r"\b(?:très intéressé|intéressé\s+\w+|très motivé|je suis convaincu)\b", letter_text):
        warnings.append("Formulation genrée ou maladroite à corriger")
        score -= 4
    if re.search(r"\.{2,}|\?{2,}|\!{2,}", " ".join(cover_letter_paragraphs or [])):
        warnings.append("Répétition de ponctuation ou formulation maladroite")
        score -= 3

    score = max(0, min(100, round(score, 1)))

    if critical or score < 70:
        status = "CRITIQUE"
        decision = "NE PAS ENVOYER"
    elif score < 90:
        status = "À VÉRIFIER"
        decision = "RELIRE"
    else:
        status = "AMÉLIORATION"
        decision = "PRÊT À ENVOYER"

    human_checks = []
    if critical:
        human_checks.append("Vérifier les éléments critiques listés ci-dessus avant envoi.")
    if warnings:
        human_checks.append("Relire les éléments signalés et valider les chiffres, dates et mentions d’entreprise.")
    human_checks.append("Valider la cohérence entre le CV et la lettre avant le dépôt final.")

    return {
        "score": score,
        "status": status,
        "decision": decision,
        "issues": issues[:8],
        "warnings": warnings[:8],
        "critical": critical[:5],
        "human_checks": human_checks[:4],
    }

def audit_cover_letter_structure(cl: dict) -> dict:
    """Vérifie que la lettre respecte la structure Vous · Moi · Nous."""
    if not isinstance(cl, dict):
        return {"score": 0, "missing": ["structure absente"], "too_short": [], "warnings": []}

    required = ["accroche", "vous", "moi", "nous", "action", "politesse"]
    missing = []
    too_short = []
    warnings = []

    for key in required:
        val = (cl.get(key) or "").strip()
        if not val:
            missing.append(key)
        elif len(val) < 60:
            too_short.append(key)

    full_text = " ".join([str(cl.get(k, "")) for k in required]).lower()
    if re.search(r"\b(?:très intéressé|intéressé par|convaincu)\b", full_text):
        warnings.append("Formulations genrées à corriger (écrire au féminin)")

    obj = (cl.get("object") or "").lower()
    accroche = (cl.get("accroche") or "").lower()
    if obj and obj[:30] and obj[:30] in accroche:
        warnings.append("L'accroche répète l'objet — à reformuler")

    if re.search(r"\b(bonjour|cher monsieur|chère madame|cordialement|bien à vous)\b", full_text):
        warnings.append("Formule familière détectée (à éviter en France)")

    score = 100
    score -= len(missing) * 15
    score -= len(too_short) * 5
    score -= len(warnings) * 5
    score = max(0, min(100, score))

    return {
        "score": score,
        "missing": missing,
        "too_short": too_short,
        "warnings": warnings
    }

def apply_safe_autocorrection(cv_data: dict, cover_letter_paragraphs: list) -> tuple:
    replacements = {
        "très intéressé": "très intéressée",
        "intéressé par": "intéressée par",
        "intéressé à": "intéressée à",
        "je suis convaincu": "je suis convaincue",
        "très motivé": "très motivée",
        "je suis très motivé": "je suis très motivée",
        "je suis très intéressé": "je suis très intéressée",
    }

    def fix(text):
        if not isinstance(text, str):
            return text
        for bad, good in replacements.items():
            text = re.sub(re.escape(bad), good, text, flags=re.IGNORECASE)
        return text

    corrected_cv = cv_data.copy()
    if isinstance(corrected_cv.get("tailored_summary"), str):
        corrected_cv["tailored_summary"] = fix(corrected_cv["tailored_summary"])
    if isinstance(corrected_cv.get("accroche_objective"), str):
        corrected_cv["accroche_objective"] = fix(corrected_cv["accroche_objective"])

    cl = {}.get("cover_letter_struct")
    corrected_cover_struct = None
    if isinstance(cl, dict):
        corrected_cover_struct = dict(cl)
        for key in ("accroche", "vous", "moi", "nous", "action", "politesse"):
            if key in corrected_cover_struct:
                corrected_cover_struct[key] = fix(corrected_cover_struct[key])
        corrected_cover_struct["paragraphs"] = cover_letter_to_paragraphs(corrected_cover_struct)

    corrected_cover = [fix(p) for p in cover_letter_paragraphs]

    if corrected_cover_struct is not None:
        {}["cover_letter_struct"] = corrected_cover_struct
        corrected_cover = corrected_cover_struct["paragraphs"]

    return corrected_cv, corrected_cover

def set_docx_margins_compact(doc):
    for section in doc.sections:
        section.top_margin = Cm(1.2)
        section.bottom_margin = Cm(1.2)
        section.left_margin = Cm(1.2)
        section.right_margin = Cm(1.2)

def set_font_run(run, font_name=FONT_FAMILY, size_pt=9.5, bold=False, color_rgb=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:ascii"), font_name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size_pt)
    run.bold = bold
    if color_rgb:
        run.font.color.rgb = color_rgb

def build_docx_documents(cv_data: dict, cover_letter_paragraphs: list, company_name: str, location: str, accent_color_rgb):
    now = datetime.now()
    today_date = f"{now.day} {MOIS_FR[now.month - 1]} {now.year}"
    job_title = cv_data.get("cv_title", MASTER_CV["default_title"])
    contact = MASTER_CV["contact"]
    contact_info = f"{location}, France | {contact['phone']} | {contact['email']} | {contact['linkedin']}"

    # -----------------------------------------------------------------------
    # Document LM — structure classique française (Vous · Moi · Nous)
    # -----------------------------------------------------------------------
    doc_lm = Document()
    set_docx_margins_compact(doc_lm)

    cl = {}.get("cover_letter_struct") or {}
    recipient = cl.get("recipient", {}) or {}
    greeting = recipient.get("title") or "Madame, Monsieur,"

    # Bloc expéditeur
    p_sender = doc_lm.add_paragraph()
    p_sender.paragraph_format.space_after = Pt(2)
    set_font_run(p_sender.add_run(MASTER_CV["candidate_name"] + "\n"), size_pt=11, bold=True, color_rgb=accent_color_rgb)
    set_font_run(p_sender.add_run(f"{job_title}\n"), size_pt=9.5)
    set_font_run(p_sender.add_run(f"{contact['phone']} | {contact['email']}\n"), size_pt=8.5, color_rgb=RGBColor(80, 80, 80))
    set_font_run(p_sender.add_run(f"{contact['linkedin']}\n"), size_pt=8.5, color_rgb=RGBColor(80, 80, 80))
    set_font_run(p_sender.add_run(f"{location}, France"), size_pt=8.5, color_rgb=RGBColor(80, 80, 80))

    # Bloc destinataire
    recipient_lines = []
    if recipient.get("name"):
        recipient_lines.append(recipient["name"])
    if recipient.get("role"):
        recipient_lines.append(recipient["role"])
    if recipient.get("company") or company_name:
        recipient_lines.append(recipient.get("company") or company_name)
    if recipient.get("address"):
        recipient_lines.append(recipient["address"])

    if recipient_lines:
        p_dest = doc_lm.add_paragraph()
        p_dest.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_dest.paragraph_format.space_before = Pt(6)
        set_font_run(p_dest.add_run("\n".join(recipient_lines)), size_pt=9.5)

    p_date = doc_lm.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_date.paragraph_format.space_before = Pt(6)
    set_font_run(p_date.add_run(f"Fait à {location}, le {today_date}"), size_pt=9.5)

    p_obj = doc_lm.add_paragraph()
    p_obj.paragraph_format.space_before = Pt(14)
    obj_text = cl.get("object") or f"Candidature au poste de {job_title}"
    if obj_text.lower().startswith("objet"):
        obj_text = re.sub(r"^objet\s*:\s*", "", obj_text, flags=re.I)
    set_font_run(p_obj.add_run("Objet : "), size_pt=10, bold=True)
    set_font_run(p_obj.add_run(obj_text), size_pt=10, bold=True, color_rgb=accent_color_rgb)

    p_greet = doc_lm.add_paragraph()
    p_greet.paragraph_format.space_before = Pt(12)
    set_font_run(p_greet.add_run(greeting), size_pt=9.5)

    for key in ("accroche", "vous", "moi", "nous", "action"):
        val = cl.get(key, "")
        if val:
            p = doc_lm.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            set_font_run(p.add_run(val), size_pt=9.5)

    pol = cl.get("politesse", "")
    if pol:
        p = doc_lm.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_font_run(p.add_run(pol), size_pt=9.5)

    p_sig = doc_lm.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig.paragraph_format.space_before = Pt(10)
    set_font_run(p_sig.add_run(cl.get("signature") or MASTER_CV["candidate_name"]), size_pt=10, bold=True, color_rgb=accent_color_rgb)

    if not cl:
        for p_text in cover_letter_paragraphs:
            p = doc_lm.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            set_font_run(p.add_run(p_text), size_pt=9.5)

    # -----------------------------------------------------------------------
    # Document CV
    # -----------------------------------------------------------------------
    doc_cv = Document()
    set_docx_margins_compact(doc_cv)

    cv_style_selected = {}.get("cv_style", "Chronologique (classique)")
    seo_title = cv_data.get("seo_title") or cv_data.get("cv_title", "")

    p_cv_header = doc_cv.add_paragraph()
    p_cv_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cv_header.paragraph_format.space_after = Pt(2)
    set_font_run(p_cv_header.add_run(MASTER_CV["candidate_name"] + "\n"), size_pt=13, bold=True, color_rgb=accent_color_rgb)
    set_font_run(p_cv_header.add_run(f"{seo_title.upper()}\n"), size_pt=10, bold=True, color_rgb=accent_color_rgb)
    set_font_run(p_cv_header.add_run(f"{contact_info}\n"), size_pt=8.5, color_rgb=RGBColor(80, 80, 80))

    # Accroche (Objectif + Compétences clés)
    accroche = cv_data.get("accroche_objective") or cv_data.get("tailored_summary", "")
    if accroche:
        p_acc = doc_cv.add_paragraph()
        p_acc.paragraph_format.space_before = Pt(6)
        p_acc.paragraph_format.space_after = Pt(6)
        p_acc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_font_run(p_acc.add_run(accroche), size_pt=9.5, color_rgb=RGBColor(60, 60, 60))

    # Compétences
    p_sec_sk = doc_cv.add_paragraph()
    p_sec_sk.paragraph_format.space_before = Pt(4)
    p_sec_sk.paragraph_format.space_after = Pt(1)
    set_font_run(p_sec_sk.add_run("COMPÉTENCES"), size_pt=10.5, bold=True, color_rgb=accent_color_rgb)

    skills_str = ", ".join(cv_data.get("skills", [])) if isinstance(cv_data.get("skills"), list) else cv_data.get("skills", "")
    tools_str = ", ".join(cv_data.get("tools", [])) if isinstance(cv_data.get("tools"), list) else cv_data.get("tools", "")

    p_sk1 = doc_cv.add_paragraph()
    p_sk1.paragraph_format.left_indent = Cm(0.3)
    p_sk1.paragraph_format.space_after = Pt(1)
    set_font_run(p_sk1.add_run("• Compétences clés : "), size_pt=9.5, bold=True)
    set_font_run(p_sk1.add_run(skills_str), size_pt=9.5)

    p_sk2 = doc_cv.add_paragraph()
    p_sk2.paragraph_format.left_indent = Cm(0.3)
    p_sk2.paragraph_format.space_after = Pt(4)
    set_font_run(p_sk2.add_run("• Outils & Technologies : "), size_pt=9.5, bold=True)
    set_font_run(p_sk2.add_run(tools_str), size_pt=9.5)

    # Modèle Thématique : Compétences & Réalisations
    if "Thématique" in cv_style_selected:
        grouped = cv_data.get("skills_grouped", []) or []
        if grouped:
            p_sec_th = doc_cv.add_paragraph()
            p_sec_th.paragraph_format.space_before = Pt(6)
            p_sec_th.paragraph_format.space_after = Pt(1)
            set_font_run(p_sec_th.add_run("COMPÉTENCES & RÉALISATIONS"), size_pt=10.5, bold=True, color_rgb=accent_color_rgb)

            for group in grouped:
                p_th = doc_cv.add_paragraph()
                p_th.paragraph_format.left_indent = Cm(0.3)
                p_th.paragraph_format.space_before = Pt(2)
                p_th.paragraph_format.space_after = Pt(0.5)
                set_font_run(p_th.add_run(str(group.get("theme", "")) + " :"), size_pt=9.5, bold=True, color_rgb=accent_color_rgb)

                for ev in group.get("evidence", []):
                    p_ev = doc_cv.add_paragraph()
                    p_ev.paragraph_format.left_indent = Cm(0.6)
                    p_ev.paragraph_format.space_after = Pt(1)
                    set_font_run(p_ev.add_run("– "), size_pt=9.5, bold=True, color_rgb=accent_color_rgb)
                    set_font_run(p_ev.add_run(ev), size_pt=9.5)

    # Expériences
    p_sec2 = doc_cv.add_paragraph()
    p_sec2.paragraph_format.space_before = Pt(6)
    p_sec2.paragraph_format.space_after = Pt(1)
    set_font_run(p_sec2.add_run("EXPÉRIENCES PROFESSIONNELLES"), size_pt=10.5, bold=True, color_rgb=accent_color_rgb)

    for exp in cv_data.get("experiences", []):
        p_head = doc_cv.add_paragraph()
        p_head.paragraph_format.space_before = Pt(3)
        p_head.paragraph_format.space_after = Pt(0.5)
        set_font_run(p_head.add_run(exp.get("role", "")), size_pt=9.5, bold=True)
        set_font_run(p_head.add_run(f" — {exp.get('company', '')} ({exp.get('period', '')})"), size_pt=9.5)

        if "Mixte" in cv_style_selected and exp.get("skill_subheadings"):
            for sub in exp.get("skill_subheadings", []):
                p_sub = doc_cv.add_paragraph()
                p_sub.paragraph_format.left_indent = Cm(0.3)
                p_sub.paragraph_format.space_before = Pt(1)
                p_sub.paragraph_format.space_after = Pt(0.5)
                set_font_run(p_sub.add_run(str(sub.get("label", "")) + " "), size_pt=9.5, bold=True, color_rgb=accent_color_rgb)
                set_font_run(p_sub.add_run(sub.get("text", "")), size_pt=9.5)
        else:
            bullets = exp.get("bullets", [])
            if isinstance(bullets, str):
                bullets = [b.strip() for b in bullets.split("\n") if b.strip()]
            for bullet in bullets:
                p_b = doc_cv.add_paragraph()
                p_b.paragraph_format.left_indent = Cm(0.3)
                p_b.paragraph_format.space_after = Pt(1)
                set_font_run(p_b.add_run("• "), size_pt=9.5, bold=True, color_rgb=accent_color_rgb)
                set_font_run(p_b.add_run(bullet), size_pt=9.5)

    # Bénévolat
    p_sec_vol = doc_cv.add_paragraph()
    p_sec_vol.paragraph_format.space_before = Pt(4)
    p_sec_vol.paragraph_format.space_after = Pt(1)
    set_font_run(p_sec_vol.add_run("ENGAGEMENT BÉNÉVOLE"), size_pt=10.5, bold=True, color_rgb=accent_color_rgb)

    for vol in cv_data.get("volunteering", []):
        p_head_v = doc_cv.add_paragraph()
        p_head_v.paragraph_format.space_before = Pt(3)
        p_head_v.paragraph_format.space_after = Pt(0.5)
        set_font_run(p_head_v.add_run(vol.get("role", "")), size_pt=9.5, bold=True)
        set_font_run(p_head_v.add_run(f" — {vol.get('organization', '')} ({vol.get('period', '')})"), size_pt=9.5)

        bullets = vol.get("bullets", [])
        if isinstance(bullets, str):
            bullets = [b.strip() for b in bullets.split("\n") if b.strip()]
        for bullet in bullets:
            p_bv = doc_cv.add_paragraph()
            p_bv.paragraph_format.left_indent = Cm(0.3)
            p_bv.paragraph_format.space_after = Pt(1)
            set_font_run(p_bv.add_run("• "), size_pt=9.5, bold=True, color_rgb=accent_color_rgb)
            set_font_run(p_bv.add_run(bullet), size_pt=9.5)

    # Formations
    p_sec3 = doc_cv.add_paragraph()
    p_sec3.paragraph_format.space_before = Pt(4)
    p_sec3.paragraph_format.space_after = Pt(1)
    set_font_run(p_sec3.add_run("FORMATIONS ET DIPLÔMES"), size_pt=10.5, bold=True, color_rgb=accent_color_rgb)

    for edu in cv_data.get("education", []):
        p_f = doc_cv.add_paragraph()
        p_f.paragraph_format.left_indent = Cm(0.3)
        p_f.paragraph_format.space_after = Pt(1)
        set_font_run(p_f.add_run("• "), size_pt=9.5, bold=True, color_rgb=accent_color_rgb)
        set_font_run(p_f.add_run(edu.get("title", "")), size_pt=9.5, bold=True)
        set_font_run(p_f.add_run(f" — {edu.get('institution', '')} ({edu.get('year', '')})"), size_pt=9.5)

    # Langues
    p_sec4 = doc_cv.add_paragraph()
    p_sec4.paragraph_format.space_before = Pt(4)
    p_sec4.paragraph_format.space_after = Pt(1)
    set_font_run(p_sec4.add_run("LANGUES"), size_pt=10.5, bold=True, color_rgb=accent_color_rgb)

    languages = cv_data.get("languages", [])
    if isinstance(languages, str):
        languages = [l.strip() for l in languages.split("\n") if l.strip()]
    for lang in languages:
        p_l = doc_cv.add_paragraph()
        p_l.paragraph_format.left_indent = Cm(0.3)
        p_l.paragraph_format.space_after = Pt(1)
        set_font_run(p_l.add_run("• "), size_pt=9.5, bold=True, color_rgb=accent_color_rgb)
        set_font_run(p_l.add_run(lang), size_pt=9.5)

    # Mais aussi…
    interests = cv_data.get("interests", []) or MASTER_CV.get("interests", [])
    if interests:
        p_sec5 = doc_cv.add_paragraph()
        p_sec5.paragraph_format.space_before = Pt(4)
        p_sec5.paragraph_format.space_after = Pt(1)
        set_font_run(p_sec5.add_run("MAIS AUSSI…"), size_pt=10.5, bold=True, color_rgb=accent_color_rgb)
        for it in interests:
            p_i = doc_cv.add_paragraph()
            p_i.paragraph_format.left_indent = Cm(0.3)
            p_i.paragraph_format.space_after = Pt(1)
            set_font_run(p_i.add_run("• "), size_pt=9.5, bold=True, color_rgb=accent_color_rgb)
            set_font_run(p_i.add_run(it), size_pt=9.5)

    return doc_cv, doc_lm

def build_cover_letter_pdf_bytes(cl: dict, company_name: str, location: str, accent_hex: str) -> bytes:
    """Génère un PDF simple de la lettre pour aperçu."""
    if not REPORTLAB_OK:
        raise ImportError("Le package 'reportlab' doit être installé (`pip install reportlab`).")
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        leftMargin=2*cm, rightMargin=2*cm, topMargin=1.8*cm, bottomMargin=1.8*cm
    )
    styles = getSampleStyleSheet()
    body_style = ParagraphStyle(
        "body", parent=styles["Normal"], fontName="Helvetica",
        fontSize=10, leading=14, alignment=TA_JUSTIFY, spaceAfter=8
    )
    right_style = ParagraphStyle("right", parent=body_style, alignment=TA_RIGHT)
    accent_color = HexColor(accent_hex) if accent_hex else HexColor("#4169E1")
    bold_accent = ParagraphStyle(
        "accent", parent=styles["Normal"], fontName="Helvetica-Bold",
        fontSize=11, textColor=accent_color
    )

    story = []
    story.append(Paragraph("<b>Candidate</b>", bold_accent))
    story.append(Paragraph("Ingénieure PMO / Cheffe de Projet IT", body_style))
    story.append(Paragraph("Île-de-France, France", body_style))
    story.append(Spacer(1, 12))

    recipient = cl.get("recipient", {}) or {}
    r_lines = []
    if recipient.get("name"): r_lines.append(recipient["name"])
    if recipient.get("role"): r_lines.append(recipient["role"])
    if recipient.get("company") or company_name:
        r_lines.append(recipient.get("company") or company_name)
    if recipient.get("address"): r_lines.append(recipient["address"])
    for line in r_lines:
        story.append(Paragraph(line, right_style))

    story.append(Spacer(1, 8))
    story.append(Paragraph(f"Fait à {location}", right_style))
    story.append(Spacer(1, 18))

    obj = cl.get("object") or "Candidature au poste de ..."
    story.append(Paragraph(f"<b>Objet : {obj}</b>", body_style))
    story.append(Spacer(1, 12))
    story.append(Paragraph(recipient.get("title") or "Madame, Monsieur,", body_style))

    for key in ("accroche", "vous", "moi", "nous", "action", "politesse"):
        val = cl.get(key)
        if val:
            story.append(Paragraph(val, body_style))

    story.append(Spacer(1, 20))
    story.append(Paragraph("<b>Candidate</b>", right_style))

    doc.build(story)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes

def save_candidature_locally(cv_data: dict, cover_letter_paragraphs: list, company_name: str, location: str, accent_color_rgb):
    clean_company = re.sub(r'[\\/*?:"<>|]', "", company_name).strip() or "Entreprise"
    folder_path = os.path.join(os.getcwd(), "Candidatures", clean_company)
    os.makedirs(folder_path, exist_ok=True)

    doc_cv, doc_lm = build_docx_documents(cv_data, cover_letter_paragraphs, clean_company, location, accent_color_rgb)

    cv_path = os.path.join(folder_path, f"CV_Ornella_SOH_{clean_company}.docx")
    lm_path = os.path.join(folder_path, f"LM_Ornella_SOH_{clean_company}.docx")

    doc_cv.save(cv_path)
    doc_lm.save(lm_path)

    return folder_path, cv_path, lm_path
